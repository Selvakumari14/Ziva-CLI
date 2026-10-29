#!/usr/bin/env python3
"""
Timon Data Processor
────────────────────
1. Import JSON → Parquet (one file per top-level key)
2. Ask user for timezone → convert date fields
3. Save timezone-corrected Parquet files
4. Run sleep processing pipeline (Phase 1 + 2 + 3) if needed
5. Save sleep_processed_timezone_corrected.parquet / .json
"""

import sys
import json
import math
import os
import argparse
from datetime import datetime, timezone, timedelta
from pathlib import Path

try:
    import pyarrow as pa
    import pyarrow.parquet as pq
    import pandas as pd
    import pytz
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Installing required packages...")
    os.system(f"{sys.executable} -m pip install pyarrow pandas pytz openpyxl --break-system-packages -q")
    import pyarrow as pa
    import pyarrow.parquet as pq
    import pandas as pd
    import pytz
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# Constants (mirror sleepCorrectionLayer.ts v3.1)
# ─────────────────────────────────────────────
HR_MIN              = 40
HR_MAX              = 120
HR_DROP_PCT         = 0.08
HRV_SLEEP_BOOST     = 1.15
SPO2_SLEEP_MAX      = 94
SPO2_AWAKE_MIN      = 96
TEMP_RISE_SLEEP     = 0.3
ACTIVITY_WAKE_STEPS = 30
ONSET_VOTES_NEEDED  = 2
WAKE_VOTES_NEEDED   = 2
FETCH_BEFORE        = 21600   # 6h before window for daytime HR baseline
FETCH_AFTER         = 1800
EXTEND_AFTER        = 7200    # max seconds to extend past firmware end
EXTEND_BEFORE       = 7200    # max seconds to extend before firmware start
LONG_GAP_MIN        = 600
LONG_GAP_MAX        = 28800
LONG_GAP_HR_FRAC    = 0.40
STITCH_GAP_SHORT    = 120
STITCH_GAP_MED      = 300
STITCH_GAP_LONG     = 1200
CONF_FIRMWARE       = 0.90
CONF_ENVELOPE       = 0.50
CONF_GAP            = 0.40
CONF_BOUT           = 0.70
BOUT_MAX_LEN        = 1
BLOCK_SEP           = 1800
RING_OFF_MARGIN     = 2.0

# Morning over-extension guard (Phase 3, Step 5.5 — trailing-synthetic-tail trim).
# A symmetric leading-synthetic-head trim (Step 5.5b, evening onset guard) also
# runs in classify_blocks; it needs no extra constants.
# When Phase-1 wake detection fails, eff_end falls back to last_fw + EXTEND_AFTER and
# Phase 2 fills that span with synthetic envelope/gap minutes, pushing the primary END
# past a plausible wake. Only trim when the primary END lands in this late-morning band.
MORNING_TRIM_FLOOR_MIN = 7 * 60    # 07:00 local — earliest end we'll consider trimming
MORNING_TRIM_CEIL_MIN  = 14 * 60   # 14:00 local — upper guard so evening blocks are never touched

# HR + movement END refinement (Phase 3, Step 5.6). Pulls the primary END back to
# the last minute that is both HR-at-sleep-level AND movement-quiet, but only when
# actual stepping afterwards confirms the person got up.
END_BASE_PCTILE      = 0.20   # personal sleep-HR baseline = 20th pct of in-block HR
END_WAKE_MARGIN_BPM  = 12     # wake threshold = baseline + max(this, frac*baseline)
END_WAKE_MARGIN_FRAC = 0.18
END_QUIET_STEPS      = 80     # steps within ±10min for a sample to still count "asleep"
END_WAKE_STEPS       = 150    # steps after the cut required to confirm wake (movement)
END_MIN_TRIM         = 1800   # 30min — only correct clear over-extensions
END_MIN_SLEEP        = 7200   # 2h — never collapse the primary block below this
END_WAKE_SUSTAIN     = 1200   # 20min — the awake run after the cut must persist
END_GAP_TOL          = 900    # 15min — tolerance for step/HR sample proximity

# Each entry: (display_name, utc_offset_hours, label_str, pytz_object)
TIMEZONES = {
    "1": ("UTC",      0.0,   "UTC+0:00",  pytz.utc),
    "2": ("IST",      +5.5,  "UTC+5:30",  pytz.timezone("Asia/Kolkata")),
    "3": ("PST",      -8.0,  "UTC-8:00",  pytz.timezone("US/Pacific")),
    "4": ("Malaysia", +8.0,  "UTC+8:00",  pytz.timezone("Asia/Kuala_Lumpur")),
    "5": ("Mountain", -7.0,  "UTC-7:00",  pytz.timezone("US/Mountain")),
    "6": ("PDT",      -7.0,  "UTC-7:00",  pytz.timezone("US/Pacific")),
}

# Health report output filenames
REPORT_FILE = "ziva_health_report_final.xlsx"
GAPS_FILE   = "ziva_health_gaps_report.xlsx"

# Health report output filenames (Processed variant — Sleep block sourced from sleep_processed)
REPORT_FILE_PROCESSED = "ziva_health_report_processed.xlsx"
GAPS_FILE_PROCESSED   = "ziva_health_gaps_report_processed.xlsx"

SLEEP_STAGE_MAP = {
    1: "Deep Sleep", 2: "Light Sleep", 3: "REM", 5: "Awake", 11: "Awake", 12: "Awake"
}

STYLE_MAP = {
    "SLEEP": {"color": "ADD8E6", "main_cols": (1, 4),  "gap_cols": (1, 3)},
    "HR":    {"color": "FFCCCB", "main_cols": (8, 9),  "gap_cols": (7, 9)},
    "SPO2":  {"color": "90EE90", "main_cols": (13, 14),"gap_cols": (12, 14)},
    "HRV":   {"color": "E6E6FA", "main_cols": (18, 23),"gap_cols": (17, 19)},
    "TEMP":  {"color": "FFDAB9", "main_cols": (26, 27),"gap_cols": (22, 24)},
}

# Variant of STYLE_MAP used when the main Report sheet needs BOTH the raw "Sleep"
# block and a second "Sleep Processed" block (Stage / Start Time / End Time /
# Duration (min)) placed right after it, before the other metrics. Everything
# from HR onward is shifted right by the width of the extra block + its gap so
# nothing overlaps. Only used for the main Report sheet — the Gaps report keeps
# using the original STYLE_MAP since sleep_processed gaps aren't tracked.
STYLE_MAP_WITH_SLEEP_PROCESSED = {
    "SLEEP":           {"color": "ADD8E6", "main_cols": (1, 4)},
    "SLEEP_PROCESSED": {"color": "87CEEB", "main_cols": (8, 11)},
    "HR":              {"color": "FFCCCB", "main_cols": (15, 16)},
    "SPO2":            {"color": "90EE90", "main_cols": (20, 21)},
    "HRV":             {"color": "E6E6FA", "main_cols": (25, 30)},
    "TEMP":            {"color": "FFDAB9", "main_cols": (33, 34)},
}

OUTPUT_DIR    = Path("./output")
CORRECTED_DIR = Path("./output_timezone_corrected")


# ─────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────
def parse_date_str(d: str) -> int:
    """'2026.04.18 16:52:01' → UTC epoch seconds"""
    if not d:
        return 0
    try:
        dt = datetime.strptime(d, "%Y.%m.%d %H:%M:%S").replace(tzinfo=timezone.utc)
        return int(dt.timestamp())
    except Exception:
        return 0

def epoch_to_date_str(epoch: int) -> str:
    """UTC epoch seconds → '2026.04.18 16:52:01'"""
    return datetime.utcfromtimestamp(epoch).strftime("%Y.%m.%d %H:%M:%S")

def percentile(vals: list, p: float):
    if not vals:
        return None
    s = sorted(vals)
    idx = (len(s) - 1) * p
    lo, hi = int(idx), math.ceil(idx)
    return s[lo] + (s[hi] - s[lo]) * (idx - lo)

def local_hour(epoch: int, tz_offset_sec: int) -> int:
    return datetime.utcfromtimestamp(epoch + tz_offset_sec).hour


# ─────────────────────────────────────────────
# STEP 1 — Import JSON → Parquet
# ─────────────────────────────────────────────
def import_json(input_file: str):
    print(f"\n📦 Step 1: Importing {input_file}...")
    print("─" * 45)
    OUTPUT_DIR.mkdir(exist_ok=True)

    with open(input_file, "r") as f:
        data = json.load(f)

    if not isinstance(data, dict):
        print("❌ Expected a top-level JSON object with table keys")
        sys.exit(1)

    for key, rows in data.items():
        if not isinstance(rows, list) or not rows:
            print(f"  ⚠️  Skipping '{key}': not a non-empty array")
            continue

        # Flatten arrays to JSON strings so Arrow can handle them
        flat_rows = []
        for row in rows:
            flat = {}
            for k, v in row.items():
                flat[k] = json.dumps(v) if isinstance(v, (list, dict)) else v
            flat_rows.append(flat)

        df = pd.DataFrame(flat_rows)
        if "date" in df.columns:
            df = df.sort_values("date").reset_index(drop=True)
        out_path = OUTPUT_DIR / f"{key}.parquet"
        df.to_parquet(out_path, index=False)
        print(f"  ✅ {key}.parquet ({len(df)} rows)")

    print(f"✅ Import complete → {OUTPUT_DIR}/")


# ─────────────────────────────────────────────
# STEP 2 — Ask timezone
# ─────────────────────────────────────────────
def ask_timezone():
    print("\n🌍 Step 2: Select your timezone for date conversion:")
    print("─" * 50)
    for k, (name, _, label, _tz) in TIMEZONES.items():
        print(f"  {k}) {name:<10} ({label})")
    print()
    valid = "/".join(TIMEZONES.keys())
    while True:
        choice = input(f"Enter your choice [{valid}]: ").strip()
        if choice in TIMEZONES:
            break
        print("❌ Invalid choice — please try again.")
    name, offset_hours, label, tz_obj = TIMEZONES[choice]
    print(f"\n✅ Selected: {name} ({label})")
    return name, offset_hours, label, tz_obj


# ─────────────────────────────────────────────
# STEP 3 — Convert dates + save corrected files
# ─────────────────────────────────────────────
def convert_dates(tz_name: str, offset_hours: float):
    print(f"\n🔄 Step 3: Converting dates to {tz_name}...")
    print("─" * 50)
    CORRECTED_DIR.mkdir(exist_ok=True)
    delta = timedelta(hours=offset_hours)

    processed = 0
    skipped   = 0

    for parquet_file in sorted(OUTPUT_DIR.glob("*.parquet")):
        table_name = parquet_file.stem
        out_path   = CORRECTED_DIR / f"{table_name}_timezone_corrected.parquet"
        print(f"\n  📄 {table_name}")

        df = pd.read_parquet(parquet_file)

        if "date" not in df.columns:
            print(f"     ℹ️  No 'date' column — copying as-is")
            df.to_parquet(out_path, index=False)
            skipped += 1
            continue

        def convert_date(d):
            try:
                dt = datetime.strptime(str(d), "%Y.%m.%d %H:%M:%S")
                return (dt + delta).strftime("%Y.%m.%d %H:%M:%S")
            except Exception:
                return d

        df["date"] = df["date"].apply(convert_date)

        # Also convert 'start' column if present (sleep table)
        if "start" in df.columns:
            df["start"] = df["start"].apply(convert_date)

        df = df.sort_values("date").reset_index(drop=True)
        df.to_parquet(out_path, index=False)
        print(f"     ✅ Saved ({len(df)} rows) → {out_path.name}")

        if table_name == "sleep":
            json_path = CORRECTED_DIR / "sleep_timezone_corrected.json"
            df.to_json(json_path, orient="records", indent=2)
            print(f"     ✅ Saved ({len(df)} rows) → sleep_timezone_corrected.json")

        processed += 1

    print(f"\n  Tables converted : {processed}")
    print(f"  Tables copied    : {skipped}")
    print(f"  Output           : {CORRECTED_DIR}/")


# ─────────────────────────────────────────────
# STEP 4 — Sleep processing pipeline
# ─────────────────────────────────────────────

def load_table(name: str) -> list[dict]:
    """Load a timezone-corrected parquet table as list of dicts."""
    path = CORRECTED_DIR / f"{name}_timezone_corrected.parquet"
    if not path.exists():
        return []
    df = pd.read_parquet(path)
    return df.to_dict("records")

def compute_baselines(heartrate, sleep_raw, hrv_table):
    cutoff = int(datetime.now(timezone.utc).timestamp()) - 30 * 86400

    hr_during_sleep = []
    for h in heartrate:
        epoch = parse_date_str(str(h.get("date", "")))
        hr    = float(h.get("singleHR", 0))
        if epoch < cutoff or not (HR_MIN <= hr <= HR_MAX):
            continue
        for s in sleep_raw:
            if int(s.get("quality", 0)) in (1, 2, 3) and abs(parse_date_str(str(s.get("date", ""))) - epoch) <= 30:
                hr_during_sleep.append(hr)
                break

    # Fallbacks mirror computeGlobalBaselinesQuery: COALESCE(..., HR_MAX)
    p90 = percentile(sorted(hr_during_sleep), 0.90) or float(HR_MAX)
    p75 = percentile(sorted(hr_during_sleep), 0.75) or float(HR_MAX)

    hrv_vals = [float(h.get("hrv", 0)) for h in hrv_table
                if h.get("hrv") and float(h.get("hrv", 0)) > 0
                and parse_date_str(str(h.get("date", ""))) >= cutoff]
    avg_hrv = sum(hrv_vals) / len(hrv_vals) if hrv_vals else 30.0

    # Detect HR interval
    epochs = sorted(parse_date_str(str(h.get("date", ""))) for h in heartrate
                    if HR_MIN <= float(h.get("singleHR", 0)) <= HR_MAX)
    gaps = [epochs[i+1] - epochs[i] for i in range(len(epochs)-1) if 61 <= epochs[i+1] - epochs[i] <= 3599]
    med_gap = percentile(sorted(gaps), 0.5) or 300
    hr_interval = 5 if med_gap <= 360 else (10 if med_gap <= 900 else 30)

    return p90, p75, avg_hrv, hr_interval


def compute_night_epochs(night_date: str, tz_offset_hours: float):
    y, m, d  = map(int, night_date.split("-"))
    midnight = int(datetime(y, m, d, tzinfo=timezone.utc).timestamp())
    prior    = midnight - 86400
    tz_sec   = round(tz_offset_hours * 3600)
    return prior + 18 * 3600 - tz_sec, midnight + 18 * 3600 - tz_sec


def phase1_signal_boundaries(window_start, window_end, heartrate, hrv_table, spo2,
                              temperature_table, activitydetails, sleep_raw, p90, avg_hrv):
    fetch_start = window_start - FETCH_BEFORE
    fetch_end   = window_end   + FETCH_AFTER
    bucket_size = 300

    def bkt(e): return (e // bucket_size) * bucket_size

    # HR buckets
    hr_map = {}
    for h in heartrate:
        e  = parse_date_str(str(h.get("date", "")))
        hr = float(h.get("singleHR", 0))
        if not (fetch_start <= e <= fetch_end) or not (HR_MIN <= hr <= HR_MAX):
            continue
        b = bkt(e)
        hr_map.setdefault(b, []).append(hr)
    hr_buckets = {b: sum(v)/len(v) for b, v in hr_map.items()}

    # Daytime HR (before window) — falls back to hrv_table.heartRate, then 75.0
    # (mirrors daytime_hr_from_hr → daytime_hr_from_hrv → 75.0 in TS v3.1)
    daytime = [v for b, v in hr_buckets.items()
               if window_start - FETCH_BEFORE <= b < window_start]
    if daytime:
        awake_hr = sum(daytime) / len(daytime)
    else:
        hrv_daytime = [
            float(h.get("heartRate", 0) or 0) for h in hrv_table
            if window_start - FETCH_BEFORE <= parse_date_str(str(h.get("date", ""))) < window_start
            and h.get("heartRate") is not None
            and HR_MIN <= float(h.get("heartRate", 0) or 0) <= HR_MAX
        ]
        awake_hr = sum(hrv_daytime) / len(hrv_daytime) if hrv_daytime else 75.0
    onset_hr = awake_hr * (1.0 - HR_DROP_PCT)
    wake_hr  = awake_hr * 0.92

    # Window HR
    window_hr = sorted([(b, v) for b, v in hr_buckets.items()
                        if window_start <= b <= window_end])
    has_hr = bool(window_hr)

    # HRV fallback
    if not has_hr:
        hrv_hr_map = {}
        for h in hrv_table:
            e  = parse_date_str(str(h.get("date", "")))
            hr = float(h.get("heartRate", 0) or 0)
            if not (window_start <= e <= window_end) or not (HR_MIN <= hr <= HR_MAX):
                continue
            hrv_hr_map.setdefault(bkt(e), []).append(hr)
        window_hr = sorted([(b, sum(v)/len(v)) for b, v in hrv_hr_map.items()])

    unified = [(b, avg, 'hrv_table' if not has_hr else 'heartrate')
               for b, avg in window_hr]

    # Activity buckets
    act_map = {}
    for a in activitydetails:
        e = parse_date_str(str(a.get("date", "")))
        if window_start <= e <= window_end:
            act_map[bkt(e)] = act_map.get(bkt(e), 0) + float(a.get("step", 0) or 0)

    # HR signal
    hr_signal = {}
    for i, (b, avg, src) in enumerate(unified):
        steps = act_map.get(b, 0)
        if src == 'hrv_table':
            v_onset = 1 if avg <= onset_hr and steps < ACTIVITY_WAKE_STEPS else 0
        else:
            n1 = unified[i+1][1] if i+1 < len(unified) else avg
            n2 = unified[i+2][1] if i+2 < len(unified) else avg
            v_onset = 1 if avg <= onset_hr and n1 <= onset_hr and n2 <= onset_hr and steps < ACTIVITY_WAKE_STEPS else 0
        hr_signal[b] = {"onset": v_onset, "wake": 1 if avg >= wake_hr else 0}

    # HRV signal
    hrv_map2 = {}
    for h in hrv_table:
        e   = parse_date_str(str(h.get("date", "")))
        hrv = float(h.get("hrv", 0) or 0)
        if window_start <= e <= window_end and hrv > 0:
            hrv_map2.setdefault(bkt(e), []).append(hrv)
    hrv_signal = {}
    thr = avg_hrv * HRV_SLEEP_BOOST
    for b, vals in hrv_map2.items():
        avg2 = sum(vals) / len(vals)
        hrv_signal[b] = {"onset": 1 if avg2 >= thr else 0, "wake": 1 if avg2 < avg_hrv else 0}

    # SpO2 signal
    spo2_map = {}
    for s in spo2:
        e   = parse_date_str(str(s.get("date", "")))
        val = float(s.get("automaticSpo2Data", 0) or 0)
        if window_start <= e <= window_end and 70 <= val <= 100:
            spo2_map.setdefault(bkt(e), []).append(val)
    spo2_signal = {}
    for b, vals in spo2_map.items():
        avg2 = sum(vals) / len(vals)
        spo2_signal[b] = {"onset": 1 if avg2 <= SPO2_SLEEP_MAX else 0, "wake": 1 if avg2 >= SPO2_AWAKE_MIN else 0}

    # Temp signal
    temp_map = {}
    for t in temperature_table:
        e    = parse_date_str(str(t.get("date", "")))
        temp = float(t.get("temperature", 0) or 0)
        if window_start <= e <= window_end and 20 <= temp <= 42:
            temp_map.setdefault(bkt(e), []).append(temp)
    temp_buckets = sorted((b, sum(v)/len(v)) for b, v in temp_map.items())
    temp_signal  = {}
    for i, (b, avg2) in enumerate(temp_buckets):
        prev_avg = temp_buckets[i-1][1] if i > 0 else avg2
        temp_signal[b] = {"onset": 1 if avg2 - prev_avg >= TEMP_RISE_SLEEP else 0, "wake": 0}

    # Activity signal
    act_signal = {b: {"onset": 1 if s < ACTIVITY_WAKE_STEPS else 0,
                      "wake":  1 if s >= ACTIVITY_WAKE_STEPS * 2 else 0}
                  for b, s in act_map.items()}

    # First sleep epoch
    first_sleep = min(
        (parse_date_str(str(s.get("date", ""))) for s in sleep_raw
         if window_start <= parse_date_str(str(s.get("date", ""))) < window_end
         and int(s.get("quality", 0)) in (1, 2, 3)),
        default=window_start
    )

    all_buckets = sorted(set(
        list(hr_signal) + list(hrv_signal) + list(spo2_signal) +
        list(temp_signal) + list(act_signal)
    ))

    min_onset = 1 if not has_hr else 2
    onset_epoch = None
    wake_epoch  = 0

    for b in all_buckets:
        if b < first_sleep - 1800:
            continue
        onset = (hr_signal.get(b, {}).get("onset", 0) + hrv_signal.get(b, {}).get("onset", 0) +
                 spo2_signal.get(b, {}).get("onset", 0) + temp_signal.get(b, {}).get("onset", 0) +
                 act_signal.get(b, {}).get("onset", 0))
        wake  = (hr_signal.get(b, {}).get("wake", 0) + hrv_signal.get(b, {}).get("wake", 0) +
                 spo2_signal.get(b, {}).get("wake", 0) + temp_signal.get(b, {}).get("wake", 0) +
                 act_signal.get(b, {}).get("wake", 0))
        if onset_epoch is None and onset >= min_onset:
            onset_epoch = b
        if wake >= 3:
            wake_epoch = b

    return onset_epoch, wake_epoch


def phase2_correct_night(window_start, window_end, wake_epoch, heartrate, sleep_raw, p75, p90, hr_interval):
    fetch_start = window_start - FETCH_BEFORE
    fetch_end   = window_end   + FETCH_AFTER

    sleep_rows = [
        {"date": parse_date_str(str(s["date"])), "quality": int(s["quality"]),
         "fw_start": str(s.get("start", "")), "unit_length": int(s.get("unitLength", s.get("unit_length", 1)) or 1)}
        for s in sleep_raw
        if window_start <= parse_date_str(str(s.get("date", ""))) < window_end
        and int(s.get("quality", 0)) in (1, 2, 3, 5)
    ]
    if not sleep_rows:
        return []

    first_fw = min(r["date"] for r in sleep_rows)
    last_fw  = max(r["date"] for r in sleep_rows)
    eff_end  = (max(min(wake_epoch, last_fw + EXTEND_AFTER), last_fw)
                if wake_epoch > 0 else last_fw + EXTEND_AFTER)

    night_hr = [
        {"date": parse_date_str(str(h["date"])), "hr": float(h.get("singleHR", 0))}
        for h in heartrate
        if fetch_start <= parse_date_str(str(h.get("date", ""))) <= fetch_end
        and HR_MIN <= float(h.get("singleHR", 0)) <= HR_MAX
    ]

    # v3.1: firmware rows are kept from eff_start (= first_fw) onward with NO upper
    # cap — spurious late blocks are handled downstream by primary-block selection.
    fw_trimmed = [
        {**r, "confidence": CONF_FIRMWARE, "source": "firmware"}
        for r in sleep_rows if r["date"] >= first_fw
    ]

    envelope_hr = [
        h for h in night_hr
        if first_fw <= h["date"] <= eff_end and h["hr"] <= p90
        and not any(abs(s["date"] - h["date"]) <= 30 for s in sleep_rows)
    ]
    envelope_exp = []
    for h in envelope_hr:
        for v in range(hr_interval):
            d = h["date"] + v * 60
            if first_fw <= d <= eff_end:
                envelope_exp.append({"date": d, "quality": 2, "fw_start": None,
                                     "unit_length": 1, "confidence": CONF_ENVELOPE, "source": "envelope"})

    # Sessions for gap recovery
    sess = {}
    for r in sleep_rows:
        k = r["fw_start"]
        if k not in sess:
            sess[k] = {"start": r["date"], "end": r["date"]}
        sess[k]["start"] = min(sess[k]["start"], r["date"])
        sess[k]["end"]   = max(sess[k]["end"],   r["date"])
    sess_list = sorted(sess.values(), key=lambda x: x["start"])

    gap_recovered = []
    for i in range(len(sess_list) - 1):
        gs, ge = sess_list[i]["end"], sess_list[i+1]["start"]
        if not (LONG_GAP_MIN <= ge - gs <= LONG_GAP_MAX):
            continue
        gap_hr = [h for h in night_hr if gs < h["date"] < ge]
        if len(gap_hr) < 3:
            continue
        low = sum(1 for h in gap_hr if h["hr"] <= p75)
        if low / len(gap_hr) < LONG_GAP_HR_FRAC:
            continue
        for h in gap_hr:
            for v in range(hr_interval):
                d = h["date"] + v * 60
                if gs < d < ge:
                    gap_recovered.append({"date": d, "quality": 2, "fw_start": None,
                                          "unit_length": 1, "confidence": CONF_GAP, "source": "gap"})

    # Deduplicate (firmware wins)
    by_date = {}
    for r in fw_trimmed + envelope_exp + gap_recovered:
        if r["date"] not in by_date or r["source"] == "firmware":
            by_date[r["date"]] = r
    deduped = sorted(by_date.values(), key=lambda x: x["date"])

    # Stitch sessions — v3.1: the SQL CASE keeps the same session for any gap up
    # to STITCH_GAP_LONG (1200s); only gaps beyond that start a new session.
    session_id = 0
    stitched   = []
    for i, row in enumerate(deduped):
        prev = deduped[i-1] if i > 0 else None
        gap  = row["date"] - prev["date"] if prev else 0
        if prev and gap > STITCH_GAP_LONG:
            session_id += 1
        stitched.append({**row, "session_id": session_id})

    sess_starts = {}
    for r in stitched:
        sid = r["session_id"]
        if sid not in sess_starts or r["date"] < sess_starts[sid]:
            sess_starts[sid] = r["date"]

    return [
        {"date": epoch_to_date_str(r["date"]),
         "quality": r["quality"],
         "start": epoch_to_date_str(sess_starts[r["session_id"]]),
         "unitLength": 1,   # v3.1: SQL output forces unitLength = 1
         "source": r["source"],
         "confidence": r["confidence"],
         "session_id": r["session_id"]}
        for r in stitched
    ]


def classify_blocks(rows: list, tz_offset_sec: int, vitals: list | None = None) -> list:
    """Phase 3 — bout consolidation + block classification (mirrors TS classifyBlocks).

    `vitals` is an optional list of {"kind": "hr"|"step", "epoch": int, "value": float}
    samples (see build_morning_vitals). When present, Step 5.6 END refinement runs;
    when omitted, behaviour matches the legacy path.
    """
    if not rows:
        return []

    def local_h(epoch): return datetime.utcfromtimestamp(epoch + tz_offset_sec).hour

    # Step 1: bout IDs
    bout_id = 0
    for i, row in enumerate(rows):
        prev = rows[i-1] if i > 0 else None
        is_start = (not prev
                    or row["quality"] != prev["quality"]
                    or parse_date_str(row["date"]) - parse_date_str(prev["date"]) > 65
                    or row["quality"] == 5
                    or (prev and prev["quality"] == 5))
        if is_start:
            bout_id += 1
        row["bout_id"] = bout_id

    bout_lens = {}
    bout_quals = {}
    for r in rows:
        bout_lens[r["bout_id"]]  = bout_lens.get(r["bout_id"], 0) + 1
        bout_quals[r["bout_id"]] = r["quality"]

    # Step 2: consolidate short deep bouts
    for r in rows:
        bid = r["bout_id"]
        if (r["quality"] == 1 and bout_lens[bid] <= BOUT_MAX_LEN
                and (bout_quals.get(bid-1) == 2 or bout_quals.get(bid+1) == 2)):
            r["quality"]    = 2
            r["source"]     = "bout_consolidated"
            r["confidence"] = CONF_BOUT

    # Step 3: block IDs
    block_id = 0
    for i, row in enumerate(rows):
        prev = rows[i-1] if i > 0 else None
        if not prev or parse_date_str(row["date"]) - parse_date_str(prev["date"]) > BLOCK_SEP:
            block_id += 1
        row["block_id"] = block_id

    block_sleep = {}
    block_night = {}
    block_fw    = {}
    sess_sizes  = {}
    sess_fw     = {}

    for r in rows:
        bid = r["block_id"]
        sid = r["session_id"]
        if r["quality"] in (1, 2, 3):
            block_sleep[bid] = block_sleep.get(bid, 0) + 1
            h = local_h(parse_date_str(r["date"]))
            if h >= 20 or h < 10:
                block_night[bid] = block_night.get(bid, 0) + 1
        if r["source"] == "firmware":
            block_fw[bid] = block_fw.get(bid, 0) + 1
            sess_fw[sid]  = sess_fw.get(sid, 0) + 1
        sess_sizes[sid] = sess_sizes.get(sid, 0) + 1

    # Prefer the block with the most night-hour sleep; fall back to total sleep
    # minutes when no block has any night sleep (e.g. shift workers).
    primary_score = block_night if block_night else block_sleep
    primary_block = str(max(primary_score, key=lambda b: (primary_score[b], -b))) if primary_score else None

    # Step 5.5: Trailing-synthetic-tail trim (morning over-extension guard).
    # If the primary block ends in the late-morning band AND its tail is purely
    # synthetic envelope/gap fill, walk the contiguous trailing run back to the
    # last firmware-backed minute and demote those rows out of the primary block.
    trimmed_idx = set()

    def is_fw_backed(src): return src in ("firmware", "bout_consolidated")

    if primary_block is not None and block_fw.get(int(primary_block), 0) > 0:
        prim_idx = [i for i, r in enumerate(rows) if str(r["block_id"]) == primary_block]
        if prim_idx:
            last_epoch = parse_date_str(rows[prim_idx[-1]]["date"])
            end_local_min = (((last_epoch + tz_offset_sec) % 86400) + 86400) % 86400 / 60
            if MORNING_TRIM_FLOOR_MIN <= end_local_min < MORNING_TRIM_CEIL_MIN:
                for p in reversed(prim_idx):
                    if is_fw_backed(rows[p]["source"]):
                        break
                    trimmed_idx.add(p)

    # Step 5.5b: Leading-synthetic-head trim (evening onset guard).
    # Symmetric to Step 5.5: if the primary block *begins* with a purely synthetic
    # envelope/gap run (HR-inferred drowsy fill before the ring actually scored
    # sleep), walk that contiguous leading run forward to the first firmware-backed
    # minute and demote those rows out of the primary block. Envelope fill bridges
    # dropouts and extends the drowsy shoulder; it does not define onset. The walk
    # stops at the first firmware/bout_consolidated row, so a block that
    # legitimately opens on firmware sleep is never touched. The block_fw guard
    # guarantees the walk terminates before the whole block is consumed.
    if primary_block is not None and block_fw.get(int(primary_block), 0) > 0:
        prim_idx = [i for i, r in enumerate(rows) if str(r["block_id"]) == primary_block]
        for p in prim_idx:
            if is_fw_backed(rows[p]["source"]):
                break
            trimmed_idx.add(p)

    # Step 5.6: HR + movement END refinement.
    # Uses the ring's own HR and step streams to pull the primary END back to the
    # last minute that is genuinely sleep (HR at personal sleep baseline AND no
    # stepping around it), provided sustained steps afterwards confirm the person
    # got up. Movement-gated and floored at END_MIN_TRIM so already-correct nights
    # and brief pre-wake arousals are never touched.
    if vitals and primary_block is not None:
        hr = sorted(
            ({"e": int(v["epoch"]), "v": float(v["value"])}
             for v in vitals
             if v.get("kind") == "hr" and v.get("value") is not None
             and math.isfinite(float(v["value"])) and float(v["value"]) > 0),
            key=lambda x: x["e"]
        )
        steps = sorted(
            ({"e": int(v["epoch"]), "v": float(v["value"])}
             for v in vitals
             if v.get("kind") == "step" and v.get("value") is not None
             and math.isfinite(float(v["value"]))),
            key=lambda x: x["e"]
        )

        def steps_in(a, b):
            return sum(x["v"] for x in steps if a < x["e"] <= b)

        prim_idx = [i for i, r in enumerate(rows) if str(r["block_id"]) == primary_block]
        if prim_idx:
            proc_start = parse_date_str(rows[prim_idx[0]]["date"])
            proc_end   = parse_date_str(rows[prim_idx[-1]]["date"])

            in_block = sorted(h["v"] for h in hr if proc_start <= h["e"] <= proc_end)
            if len(in_block) >= 4:
                base = in_block[min(len(in_block) - 1, math.floor(END_BASE_PCTILE * len(in_block)))]
                thr  = base + max(END_WAKE_MARGIN_BPM, END_WAKE_MARGIN_FRAC * base)

                # A sample is "asleep" only if HR is at sleep level AND steps are
                # quiet around it — immunises the anchor against transient low-HR
                # dips while the person is already up and moving.
                def asleep_at(h):
                    return (h["v"] <= thr
                            and steps_in(h["e"] - END_GAP_TOL + 300,
                                         h["e"] + END_GAP_TOL - 300) <= END_QUIET_STEPS)

                tail = [h for h in hr if proc_start <= h["e"] <= proc_end + END_GAP_TOL]
                wake_at = proc_end
                for h in reversed(tail):
                    if asleep_at(h):
                        wake_at = h["e"]
                        break

                run_len     = proc_end - wake_at
                steps_after = steps_in(wake_at, proc_end + END_GAP_TOL)
                if (wake_at < proc_end
                        and run_len >= END_WAKE_SUSTAIN
                        and steps_after >= END_WAKE_STEPS
                        and run_len >= END_MIN_TRIM
                        and (wake_at - proc_start) >= END_MIN_SLEEP):
                    for i in prim_idx:
                        if parse_date_str(rows[i]["date"]) > wake_at:
                            trimmed_idx.add(i)

    result = []
    for i, r in enumerate(rows):
        bid = r["block_id"]
        sid = r["session_id"]
        if sess_sizes.get(sid, 0) < 30 and not sess_fw.get(sid):
            bt = "microsleep"
        elif not block_fw.get(bid) and block_sleep.get(bid, 0) < 30:
            bt = "microsleep"
        elif str(bid) == primary_block:
            bt = "secondary" if i in trimmed_idx else "primary"
        else:
            bt = "secondary"

        result.append({
            "date":       r["date"],
            "quality":    r["quality"],
            "start":      r["start"],
            "unitLength": r["unitLength"],
            "source":     r["source"],
            "confidence": r["confidence"],
            "block_type": bt,
        })

    return result


def build_morning_vitals(window_start: int, window_end: int,
                         heartrate: list, activitydetails: list) -> list:
    """Mirror of computeMorningVitalsQuery — HR + step samples spanning the night
    window (plus FETCH_AFTER tail) for the Step 5.6 END refinement."""
    vitals = []
    hi = window_end + FETCH_AFTER
    for h in heartrate:
        e  = parse_date_str(str(h.get("date", "")))
        hr = float(h.get("singleHR", 0) or 0)
        if window_start <= e < hi and HR_MIN <= hr <= HR_MAX:
            vitals.append({"kind": "hr", "epoch": e, "value": hr})
    for a in activitydetails:
        e = parse_date_str(str(a.get("date", "")))
        if window_start <= e < hi:
            vitals.append({"kind": "step", "epoch": e, "value": float(a.get("step", 0) or 0)})
    vitals.sort(key=lambda v: v["epoch"])
    return vitals


def run_sleep_pipeline(tz_offset_hours: float):
    print("\n🛌 Step 4: Running sleep processing pipeline...")
    print("─" * 50)

    sleep_raw         = load_table("sleep")
    heartrate         = load_table("heartrate")
    hrv_table         = load_table("hrv_table")
    spo2              = load_table("spo2")
    temperature_table = load_table("temperature_table")
    activitydetails   = load_table("activitydetails")

    print(f"  Loaded — Sleep: {len(sleep_raw)}, HR: {len(heartrate)}, "
          f"HRV: {len(hrv_table)}, SpO2: {len(spo2)}, "
          f"Temp: {len(temperature_table)}, Activity: {len(activitydetails)}")

    if not sleep_raw:
        print("  ⚠️  No sleep data — skipping")
        return

    p90, p75, avg_hrv, hr_interval = compute_baselines(heartrate, sleep_raw, hrv_table)
    print(f"  Baselines — p90: {p90:.1f}, p75: {p75:.1f}, avgHrv: {avg_hrv:.1f}, hrInterval: {hr_interval}min")

    # Unique night dates from sleep.start
    start_epochs = sorted(
        parse_date_str(str(s.get("start", "")))
        for s in sleep_raw if s.get("start")
    )
    if not start_epochs:
        print("  ⚠️  No valid start epochs — skipping")
        return

    first_day   = datetime.utcfromtimestamp(start_epochs[0]).date()
    today       = datetime.utcnow().date()
    unique_dates = []
    d = first_day
    while d <= today:
        unique_dates.append(d.strftime("%Y-%m-%d"))
        d += timedelta(days=1)

    print(f"  Processing {len(unique_dates)} nights: {unique_dates[0]} → {unique_dates[-1]}")

    tz_offset_sec = round(tz_offset_hours * 3600)
    all_rows      = []

    for date in unique_dates:
        window_start, window_end = compute_night_epochs(date, tz_offset_hours)

        # Guard mirrors hasSleepDataQuery — skip nights with no firmware rows
        # (quality 1, 2, 3, or 5) so no signal-only rows are ever generated.
        has_sleep = any(
            window_start <= parse_date_str(str(s.get("date", ""))) < window_end
            and int(s.get("quality", 0)) in (1, 2, 3, 5)
            for s in sleep_raw
        )
        if not has_sleep:
            continue

        # Phase 1
        _, wake_epoch = phase1_signal_boundaries(
            window_start, window_end, heartrate, hrv_table,
            spo2, temperature_table, activitydetails, sleep_raw, p90, avg_hrv
        )

        # Phase 2
        phase2 = phase2_correct_night(
            window_start, window_end, wake_epoch,
            heartrate, sleep_raw, p75, p90, hr_interval
        )
        if not phase2:
            print(f"  ⚠️  {date}: no corrected rows")
            continue

        # Phase 3 — vitals feed enables the Step 5.6 HR + movement END refinement
        vitals = build_morning_vitals(window_start, window_end, heartrate, activitydetails)
        final = classify_blocks(phase2, tz_offset_sec, vitals)
        print(f"  ✅ {date}: {len(final)} rows (wake: {epoch_to_date_str(wake_epoch) if wake_epoch else 'none'})")
        all_rows.extend(final)

    print(f"\n  Total sleep_processed rows: {len(all_rows)}")

    if not all_rows:
        print("  ⚠️  No rows to write")
        return

    # Sort by date before writing so both outputs are in chronological order
    all_rows.sort(key=lambda r: r["date"])

    out_path = CORRECTED_DIR / "sleep_processed_timezone_corrected.parquet"
    df = pd.DataFrame(all_rows)
    df.to_parquet(out_path, index=False)
    print(f"  ✅ Written → {out_path}")

    json_path = CORRECTED_DIR / "sleep_processed_timezone_corrected.json"
    with open(json_path, "w") as f:
        json.dump(all_rows, f, indent=2)
    print(f"  ✅ Written → {json_path}")


# ─────────────────────────────────────────────
# STEP 5 — Post-processing analysis menu
# ─────────────────────────────────────────────

def ask_analysis(has_sleep_processed: bool):
    print("\n" + "═" * 45)
    print("🔍 Analysis")
    print("═" * 45)
    print("  Which dataset would you like to analyse?")
    print()
    print("  1) Sleep (sleep_timezone_corrected)")
    if has_sleep_processed:
        print("  2) Sleep Processed (sleep_processed_timezone_corrected)")
    print("  0) Skip / Exit")
    print()
    valid = {"0", "1", "2"} if has_sleep_processed else {"0", "1"}
    while True:
        choice = input("Enter your choice: ").strip()
        if choice in valid:
            return choice
        print("  ❌ Invalid choice — please try again.")


def ask_analysis_date() -> str:
    """Prompt for a date in YYYY.MM.DD format and return it validated."""
    while True:
        raw = input("\n  Enter the date to analyse (YYYY.MM.DD): ").strip()
        try:
            datetime.strptime(raw, "%Y.%m.%d")
            return raw
        except ValueError:
            print("  ❌ Invalid format — please use YYYY.MM.DD (e.g. 2026.03.18)")


def _build_window(date_str: str):
    """
    Return (window_start_str, window_end_str) in local time for the given date.
    Window: 18:00:00 of the previous day → 17:59:59 of the selected date.
    Both files already have tz-converted dates so no offset is applied here.
    """
    chosen = datetime.strptime(date_str, "%Y.%m.%d")
    start  = datetime(chosen.year, chosen.month, chosen.day, 18, 0, 0) - timedelta(days=1)
    end    = datetime(chosen.year, chosen.month, chosen.day, 17, 59, 59)
    return start.strftime("%Y.%m.%d %H:%M:%S"), end.strftime("%Y.%m.%d %H:%M:%S")


def _run_sleep_analysis(df_all: pd.DataFrame, label: str, date_str: str,
                        window_start_str: str, window_end_str: str,
                        restrict_to_primary: bool = True):
    """Filter, then print summary + gap analysis. Shared by both sources.

    If the data has a `block_type` column (primary/secondary/microsleep — as
    produced by classifyBlocks() / run_sleep_pipeline()), the summary and
    span are computed on the `primary` block only by default, matching how
    the JS pipeline's downstream consumers treat "the night's sleep". Any
    secondary/microsleep rows that fall inside the naive 18:00→17:59:59
    window (e.g. a nap the following morning) are excluded and reported
    separately, rather than silently stretching the last-record time.
    """
    print(f"\n🛌 {label} — {date_str}")
    print("─" * 50)
    print(f"  Window (local) : {window_start_str} → {window_end_str}")

    df = df_all[
        (df_all["date"] >= window_start_str) & (df_all["date"] <= window_end_str)
    ].copy()
    df = df.sort_values("date").reset_index(drop=True)
    df["_epoch"] = df["date"].apply(lambda d: parse_date_str(str(d)))

    print(f"\n  Records in window (all block types) : {len(df)}")

    if df.empty:
        print("  ⚠️  No sleep records found in this window.")
        return

    if "block_type" in df.columns:
        bt_counts = df["block_type"].value_counts()
        print("  Block type breakdown:")
        for bt, cnt in bt_counts.items():
            print(f"    {bt:10s} : {cnt:5d} records")

        if restrict_to_primary and "primary" in bt_counts.index:
            excluded = df[df["block_type"] != "primary"]
            df = df[df["block_type"] == "primary"].reset_index(drop=True)
            if not excluded.empty:
                print(f"  ℹ️  Excluding {len(excluded)} non-primary record(s) from summary below "
                      f"(use restrict_to_primary=False to include them).")
        elif restrict_to_primary:
            print("  ⚠️  No 'primary' block found in this window — showing all block types.")

    print(f"\n  Records used for summary : {len(df)}")

    if df.empty:
        print("  ⚠️  No primary sleep records found in this window.")
        return

    # ── Summary ───────────────────────────────────────────────────────────────
    span_min = (int(df["_epoch"].iloc[-1]) - int(df["_epoch"].iloc[0])) / 60
    print(f"  First record    : {df['date'].iloc[0]}")
    print(f"  Last  record    : {df['date'].iloc[-1]}")
    print(f"  Span            : {span_min:.1f} min  ({span_min/60:.2f} h)")

    # Quality breakdown
    if "quality" in df.columns:
        quality_labels = {1: "Deep", 2: "Light", 3: "REM", 4: "Awake", 5: "Wake"}
        q_counts = df["quality"].value_counts().sort_index()
        print("\n  Quality breakdown:")
        for q, cnt in q_counts.items():
            label_q = quality_labels.get(int(q), f"Q{q}")
            print(f"    {label_q:8s} (q={q}) : {cnt:5d} records  ({cnt:.0f} min)")
        total_sleep = sum(q_counts.get(q, 0) for q in [1, 2, 3, 5])
        print(f"    {'─' * 38}")
        print(f"    {'Total':8s}       : {total_sleep:5d} records  ({total_sleep:.0f} min  /  {total_sleep/60:.2f} h)")

    # ── Gap analysis ──────────────────────────────────────────────────────────
    print("\n  Gap analysis (expected cadence: 1 record / min)")
    print("  " + "─" * 42)

    epochs     = df["_epoch"].tolist()
    gaps_found = []
    for i in range(1, len(epochs)):
        gap_sec = epochs[i] - epochs[i - 1]
        if gap_sec > 60:
            missing_min = (gap_sec - 60) / 60
            gaps_found.append((df["date"].iloc[i - 1], df["date"].iloc[i], gap_sec, missing_min))

    if not gaps_found:
        print("  ✅ No gaps found — data is continuous throughout the window.")
    else:
        print(f"  ⚠️  {len(gaps_found)} gap(s) detected:\n")
        total_missing = 0.0
        for idx, (gs, ge, gap_sec, miss) in enumerate(gaps_found, 1):
            total_missing += miss
            print(f"  [{idx:>3}] {gs}  →  {ge}")
            print(f"         Gap : {gap_sec} s  |  Missing ≈ {miss:.1f} min")
        print(f"\n  Total missing data : ≈ {total_missing:.1f} min across {len(gaps_found)} gap(s)")

    # ── Continuous sleep segments ─────────────────────────────────────────────
    print("\n  Continuous sleep segments (no gap > 1 min between records)")
    print("  " + "─" * 42)

    quality_labels = {1: "Deep", 2: "Light", 3: "REM", 4: "Awake", 5: "Wake"}

    segments = []
    seg_start_idx = 0
    for i in range(1, len(epochs)):
        if epochs[i] - epochs[i - 1] > 60:
            segments.append((seg_start_idx, i - 1))
            seg_start_idx = i
    segments.append((seg_start_idx, len(epochs) - 1))

    # Only show segments with actual sleep (q 1/2/3/5), skip single-record noise
    sleep_qualities = {1, 2, 3, 5}
    shown = 0
    for seg_idx, (si, ei) in enumerate(segments, 1):
        seg_df    = df.iloc[si:ei + 1]
        dur_min   = (int(epochs[ei]) - int(epochs[si])) / 60
        has_sleep = seg_df["quality"].apply(lambda q: int(q) in sleep_qualities).any() if "quality" in seg_df.columns else True

        if not has_sleep:
            continue

        shown += 1
        print(f"\n  Segment {shown}  |  {df['date'].iloc[si]}  →  {df['date'].iloc[ei]}")
        print(f"            Duration : {dur_min:.1f} min  ({dur_min / 60:.2f} h)  |  {ei - si + 1} records")

        if "quality" in seg_df.columns:
            q_counts_seg = seg_df["quality"].value_counts().sort_index()
            parts = []
            for q, cnt in q_counts_seg.items():
                lbl = quality_labels.get(int(q), f"Q{int(q)}")
                parts.append(f"{lbl} {cnt} min")
            print(f"            Stages   : {',  '.join(parts)}")

    if shown == 0:
        print("  ⚠️  No continuous sleep segments found.")
    else:
        print(f"\n  Total segments : {shown}")

    print()


def _run_sleep_analysis_summary_only(df_all: pd.DataFrame, label: str,
                                      window_start_str: str, window_end_str: str,
                                      restrict_to_primary: bool = True):
    """Same filtering rules as _run_sleep_analysis (window filter + primary-block
    restriction), but prints only the title plus First record / Last record / Span
    — no window/record-count/quality/gap/segment detail."""
    df = df_all[
        (df_all["date"] >= window_start_str) & (df_all["date"] <= window_end_str)
    ].copy()
    df = df.sort_values("date").reset_index(drop=True)
    df["_epoch"] = df["date"].apply(lambda d: parse_date_str(str(d)))

    print(f"\n🛌 {label}")
    print("─" * 45)

    if df.empty:
        print("  ⚠️  No sleep records found in this window.")
        return

    if "block_type" in df.columns:
        bt_counts = df["block_type"].value_counts()
        if restrict_to_primary and "primary" in bt_counts.index:
            df = df[df["block_type"] == "primary"].reset_index(drop=True)

    if df.empty:
        print("  ⚠️  No primary sleep records found in this window.")
        return

    span_min = (int(df["_epoch"].iloc[-1]) - int(df["_epoch"].iloc[0])) / 60
    print(f"  First record    : {df['date'].iloc[0]}")
    print(f"  Last  record    : {df['date'].iloc[-1]}")
    print(f"  Span            : {span_min:.1f} min  ({span_min/60:.2f} h)")
    print()


def analyse_sleep_timezone_corrected():
    """Analyse output_timezone_corrected/sleep_timezone_corrected.json."""
    date_str = ask_analysis_date()
    window_start_str, window_end_str = _build_window(date_str)

    sleep_path = CORRECTED_DIR / "sleep_timezone_corrected.json"
    if not sleep_path.exists():
        print(f"  ❌ {sleep_path.name} not found. Run the pipeline first.")
        return

    with open(sleep_path, "r") as f:
        records = json.load(f)

    if not records:
        print("  ⚠️  sleep_timezone_corrected.json is empty.")
        return

    df_all = pd.DataFrame(records)
    _run_sleep_analysis(df_all, "Sleep Timezone Corrected Analysis", date_str, window_start_str, window_end_str)


def analyse_sleep_processed_timezone_corrected():
    """Analyse output_timezone_corrected/sleep_processed_timezone_corrected.parquet."""
    date_str = ask_analysis_date()
    window_start_str, window_end_str = _build_window(date_str)

    sleep_path = CORRECTED_DIR / "sleep_processed_timezone_corrected.parquet"
    if not sleep_path.exists():
        print(f"  ❌ {sleep_path.name} not found.")
        return

    df_all = pd.read_parquet(sleep_path)
    _run_sleep_analysis(df_all, "Sleep Processed Timezone Corrected Analysis", date_str, window_start_str, window_end_str)


def _ask_fetch_sleep_date(raw_data: dict, tz_obj) -> str | None:
    """Prompt for which night to run 'Fetch Sleep Start and End Time' against.
       1) Today      → today's date (window: yesterday 6 PM → today 6 PM)
       2) Yesterday   → yesterday's date (window: day-before-yesterday 6 PM → yesterday 6 PM)
       3) Other Dates → lists every date that has sleep records (same discovery/format
                         used by the Health Report date picker) and lets the user choose.
    Returns a 'YYYY.MM.DD' date string, or None if selection couldn't be completed
    (e.g. no dates available for option 3).
    """
    print("\n  Select which night to fetch:")
    print("  " + "─" * 40)
    print("    1) Today")
    print("    2) Yesterday")
    print("    3) Other Dates")
    print()

    while True:
        choice = input("  Enter your choice [1/2/3]: ").strip()

        if choice == "1":
            return datetime.now().strftime("%Y.%m.%d")

        elif choice == "2":
            return (datetime.now() - timedelta(days=1)).strftime("%Y.%m.%d")

        elif choice == "3":
            report_dates = _hr_discover_dates(raw_data, tz_obj)
            if not report_dates:
                print("  ❌ No dates with sleep records found.")
                return None

            print("\n  Available dates (night window: 6 PM Day-1 → 6 PM Day):")
            for i, d in enumerate(report_dates, 1):
                print(f"    {i}) {d.strftime('%d-%m-%Y')}")

            while True:
                raw = input(f"\n  Select date (1–{len(report_dates)}): ").strip()
                if raw.isdigit():
                    idx = int(raw)
                    if 1 <= idx <= len(report_dates):
                        return report_dates[idx - 1].strftime("%Y.%m.%d")
                print("  ❌ Invalid choice — please try again.")

        else:
            print("  ❌ Invalid choice — please enter 1, 2, or 3.")


def fetch_sleep_start_and_end_time(raw_data: dict, tz_obj, sleep_processed_available: bool):
    """Combined workflow for menu option 4:
       0) Prompts for which night to use — Today / Yesterday / Other Dates (see
          _ask_fetch_sleep_date).
       1) Sleep Processed (sleep_processed_timezone_corrected) summary, titled 'Test Build
          Sleep Data' — prints only First record / Last record / Span.
       2) Generate Health Report (Excel) for the same date. The generated report's Report
          sheet shows the raw Sleep block (Stage / Start Time / End Time / Duration (min)),
          immediately followed by the Sleep Processed block (same four columns, built from
          sleep_processed), followed by the other metrics (HR/SPO2/HRV/TEMP) exactly as
          before. 'Sleep Start' / 'Wake Up' (first Start Time / last End Time from the raw
          SLEEP block) is then reported as 'Ring Sleep Data'.
    """
    print("\n" + "═" * 50)
    print("🩺  Fetch Sleep Start and End Time")
    print("═" * 50)

    if not sleep_processed_available:
        print("  ❌ Sleep processed data is not available — cannot run this option.")
        return

    sleep_path = CORRECTED_DIR / "sleep_processed_timezone_corrected.parquet"
    if not sleep_path.exists():
        print(f"  ❌ {sleep_path.name} not found.")
        return

    date_str = _ask_fetch_sleep_date(raw_data, tz_obj)
    if date_str is None:
        return

    # ── Part 1: Analyse Sleep → Sleep Processed ("Test Build Sleep Data") ─────
    window_start_str, window_end_str = _build_window(date_str)

    df_all = pd.read_parquet(sleep_path)
    _run_sleep_analysis_summary_only(df_all, "Test Build Sleep Data", window_start_str, window_end_str)

    # ── Part 2: Generate Health Report (Excel) → "Ring Sleep Data" ────────────
    # include_sleep_processed_block=True adds the Sleep Processed Stage/Start Time/
    # End Time/Duration (min) columns right after the raw Sleep block in the
    # generated report, ahead of the other metrics — see _hr_build_sleep_block /
    # STYLE_MAP_WITH_SLEEP_PROCESSED.
    summary = generate_health_report(raw_data, tz_obj, preselected_date_str=date_str,
                                      return_sleep_summary=True,
                                      include_sleep_processed_block=True)

    print("\n" + "═" * 45)
    print("  Ring Sleep Data")
    print("═" * 45)
    if summary and summary.get("sleep_start") and summary.get("wake_up"):
        print(f"  Sleep Start : {summary['sleep_start']}")
        print(f"  Wake Up : {summary['wake_up']}")
    else:
        print("  ⚠️  Could not determine sleep start / wake up times from the generated report.")


# ─────────────────────────────────────────────
# HEALTH REPORT (Ziva) — helpers + generator
# ─────────────────────────────────────────────

def _hr_parse_to_tz(date_str, target_tz):
    """Parse a date string (multiple formats) and return tz-aware datetime."""
    if not date_str or not isinstance(date_str, str):
        return None
    for fmt in ["%Y.%m.%d %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"]:
        try:
            dt = datetime.strptime(date_str, fmt)
            return pytz.utc.localize(dt).astimezone(target_tz)
        except ValueError:
            continue
    return None


def _hr_get_report_date(dt):
    if dt.hour >= 18:
        return (dt + timedelta(days=1)).date()
    return dt.date()


def _hr_discover_dates(data, tz_obj):
    dates = set()
    for key in ["heartrate", "sleep", "spo2"]:
        for r in data.get(key, []):
            ts = r.get("date") or r.get("start") or r.get("startTime_SleepData")
            dt = _hr_parse_to_tz(ts, tz_obj)
            if dt:
                dates.add(_hr_get_report_date(dt))
    return sorted(list(dates))


def _hr_format_dt_12h(dt):
    return dt.strftime("%d-%m-%Y %I:%M %p") if dt else ""


def _hr_format_time_12h(dt):
    return dt.strftime("%I:%M %p") if dt else ""


def _hr_find_gaps(records, start_time, end_time, threshold_mins):
    gaps = []
    if not records:
        dur = (end_time - start_time).total_seconds() / 60
        if dur > 1:
            gaps.append({"Gap Start": _hr_format_dt_12h(start_time),
                         "Gap End": _hr_format_dt_12h(end_time),
                         "Duration (min)": int(dur)})
        return gaps

    first_gap = (records[0]["dt"] - start_time).total_seconds() / 60
    if first_gap > 1:
        gaps.append({"Gap Start": _hr_format_dt_12h(start_time),
                     "Gap End": _hr_format_dt_12h(records[0]["dt"]),
                     "Duration (min)": int(first_gap)})

    for i in range(len(records) - 1):
        gap = (records[i + 1]["dt"] - records[i]["dt"]).total_seconds() / 60
        if gap > threshold_mins:
            gaps.append({"Gap Start": _hr_format_dt_12h(records[i]["dt"]),
                         "Gap End": _hr_format_dt_12h(records[i + 1]["dt"]),
                         "Duration (min)": int(gap)})

    last_gap = (end_time - records[-1]["dt"]).total_seconds() / 60
    if last_gap > 1:
        gaps.append({"Gap Start": _hr_format_dt_12h(records[-1]["dt"]),
                     "Gap End": _hr_format_dt_12h(end_time),
                     "Duration (min)": int(last_gap)})
    return gaps


def _hr_read_report_sleep_times(filename):
    """Read back the exported Excel report's 'Report' sheet SLEEP block and
    return the first Start Time and the last End Time found there."""
    try:
        wb = load_workbook(filename, data_only=True)
    except Exception as e:
        print(f"  ⚠️  Could not open {filename} to read back sleep times: {e}")
        return None

    if "Report" not in wb.sheetnames:
        return None

    sh = wb["Report"]
    sc, ec = STYLE_MAP["SLEEP"]["main_cols"]

    # Row 2 (1-based) holds the pandas-written column headers for this block.
    headers = {}
    for c in range(sc, ec + 1):
        val = sh.cell(row=2, column=c).value
        if val:
            headers[val] = c

    start_col = headers.get("Start Time")
    end_col   = headers.get("End Time")
    if not start_col or not end_col:
        return None

    first_start, last_end = None, None
    row = 3
    while True:
        v_start = sh.cell(row=row, column=start_col).value
        v_end   = sh.cell(row=row, column=end_col).value
        if v_start is None and v_end is None:
            break
        if first_start is None and v_start:
            first_start = v_start
        if v_end:
            last_end = v_end
        row += 1

    return {"sleep_start": first_start, "wake_up": last_end}


def _hr_build_sleep_block(raw_data: dict, sleep_source: str, tz_obj, report_start, report_end):
    """Build the grouped Stage/Start Time/End Time/Duration (min) rows — plus the
    matching gap list and the flat per-minute record list — for a given top-level
    sleep source key ("sleep" or "sleep_processed") over [report_start, report_end].

    This is the same grouping logic generate_health_report has always used for its
    single SLEEP block, factored out so it can be run twice (once per source) when
    a report needs to show both blocks side by side.
    """
    sleep_raw_hr = []
    for e in raw_data.get(sleep_source, []):
        qa_str = e.get("arraySleepQuality")
        if qa_str:
            dt_s = _hr_parse_to_tz(e.get("startTime_SleepData"), tz_obj)
            ul   = int(e.get("sleepUnitLength", 1))
            try:
                qa = json.loads(qa_str)
            except Exception:
                qa = []
            for i, q in enumerate(qa):
                t = dt_s + timedelta(minutes=i * ul)
                if report_start <= t <= report_end:
                    sleep_raw_hr.append({"dt": t, "stage": SLEEP_STAGE_MAP.get(q, f"S{q}"), "mins": ul})
        else:
            dt = _hr_parse_to_tz(e.get("date") or e.get("start"), tz_obj)
            if dt and report_start <= dt <= report_end:
                sleep_raw_hr.append({"dt": dt,
                                     "stage": SLEEP_STAGE_MAP.get(e.get("quality"), "S"),
                                     "mins": int(e.get("unitLength", 1))})

    sleep_raw_hr.sort(key=lambda x: x["dt"])
    grouped_sleep, sleep_gaps = [], []

    if sleep_raw_hr:
        curr = sleep_raw_hr[0]; s_t = curr["dt"]; l_t = curr["dt"]; t_m = curr["mins"]
        start_gap = (s_t - report_start).total_seconds() / 60
        if start_gap > 5:
            sleep_gaps.append({"Gap Start": _hr_format_time_12h(report_start),
                               "Gap End": _hr_format_time_12h(s_t),
                               "Duration (min)": int(start_gap)})
        for i in range(1, len(sleep_raw_hr)):
            ent = sleep_raw_hr[i]
            exp = l_t + timedelta(minutes=sleep_raw_hr[i - 1]["mins"])
            if ent["stage"] == curr["stage"] and (ent["dt"] - exp).total_seconds() <= 60:
                l_t = ent["dt"]; t_m += ent["mins"]
            else:
                grouped_sleep.append({
                    "Stage": curr["stage"],
                    "Start Time": _hr_format_time_12h(s_t),
                    "End Time": _hr_format_time_12h(l_t + timedelta(minutes=sleep_raw_hr[i - 1]["mins"])),
                    "Duration (min)": t_m
                })
                gap = (ent["dt"] - exp).total_seconds() / 60
                if gap > 5:
                    sleep_gaps.append({"Gap Start": _hr_format_time_12h(exp),
                                       "Gap End": _hr_format_time_12h(ent["dt"]),
                                       "Duration (min)": int(gap)})
                curr = ent; s_t = ent["dt"]; l_t = ent["dt"]; t_m = ent["mins"]
        grouped_sleep.append({
            "Stage": curr["stage"],
            "Start Time": _hr_format_time_12h(s_t),
            "End Time": _hr_format_time_12h(l_t + timedelta(minutes=sleep_raw_hr[-1]["mins"])),
            "Duration (min)": t_m
        })
        last_exp = l_t + timedelta(minutes=sleep_raw_hr[-1]["mins"])
        end_gap  = (report_end - last_exp).total_seconds() / 60
        if end_gap > 5:
            sleep_gaps.append({"Gap Start": _hr_format_time_12h(last_exp),
                               "Gap End": _hr_format_time_12h(report_end),
                               "Duration (min)": int(end_gap)})

    return sleep_raw_hr, grouped_sleep, sleep_gaps


def generate_health_report(raw_data: dict, tz_obj, preselected_date_str: str | None = None,
                            return_sleep_summary: bool = False, sleep_source: str = "sleep",
                            report_filename: str | None = None, gaps_filename: str | None = None,
                            include_sleep_processed_block: bool = False):
    """Generate the Ziva health Excel report for a user-selected (or pre-selected) date.

    If `preselected_date_str` (format 'YYYY.MM.DD') is given, the matching report
    date is used automatically instead of prompting the user. If no match is found,
    falls back to the interactive prompt.

    If `return_sleep_summary` is True, after export the 'Report' sheet is read back
    and {"sleep_start": ..., "wake_up": ...} (first Start Time / last End Time of the
    SLEEP block) is returned.

    `sleep_source` selects which top-level JSON key feeds the SLEEP block of the
    report — "sleep" (default, raw sleep data) or "sleep_processed" (pipeline-
    processed sleep data, used by the "Generate Processed Health Report" option).

    `report_filename` / `gaps_filename` override the default output filenames —
    used to keep the processed report separate from the standard one.

    `include_sleep_processed_block`, when True, adds a SECOND sleep block to the
    main Report sheet — Stage / Start Time / End Time / Duration (min) built from
    "sleep_processed" — placed immediately after the primary SLEEP block (built
    from `sleep_source`) and before the other metrics (HR/SPO2/HRV/TEMP), which
    are all shifted right to make room. All existing data is otherwise unchanged.
    The Gaps report is unaffected by this flag.
    """
    report_filename = report_filename or REPORT_FILE
    gaps_filename   = gaps_filename or GAPS_FILE

    print("\n" + "═" * 50)
    print("📊  Health Report Generator" + (" (Processed)" if sleep_source == "sleep_processed" else ""))
    print("═" * 50)
    print(f"  Sleep data source : {sleep_source}")

    report_dates = _hr_discover_dates(raw_data, tz_obj)
    if not report_dates:
        print("  ❌ No data found in JSON for health report.")
        return None

    print("\n  Available dates (night window: 6 PM Day-1 → 6 PM Day):")
    for i, d in enumerate(report_dates, 1):
        print(f"    {i}) {d.strftime('%d-%m-%Y')}")

    selected_date = None

    if preselected_date_str:
        try:
            wanted = datetime.strptime(preselected_date_str, "%Y.%m.%d").date()
            for d in report_dates:
                if d == wanted:
                    selected_date = d
                    break
            if selected_date is None:
                for d in report_dates:
                    if d.day == wanted.day and d.month == wanted.month:
                        selected_date = d
                        break
        except ValueError:
            selected_date = None

        if selected_date is not None:
            print(f"\n  ✅ Auto-selected date {selected_date.strftime('%d-%m-%Y')} "
                  f"(from {preselected_date_str})")
        else:
            print(f"\n  ⚠️  No report date matches {preselected_date_str} — please select manually.")

    while selected_date is None:
        raw = input(f"\n  Select date (1–{len(report_dates)}): ").strip()
        if raw.isdigit():
            idx = int(raw)
            if 1 <= idx <= len(report_dates):
                selected_date = report_dates[idx - 1]
            else:
                # Try matching by day number
                for d in report_dates:
                    if d.day == idx:
                        selected_date = d
                        break
        if selected_date is None:
            print("  ❌ Invalid choice — please try again.")

    report_start = tz_obj.localize(
        datetime.combine(selected_date - timedelta(days=1), datetime.min.time()).replace(hour=18))
    report_end = tz_obj.localize(
        datetime.combine(selected_date, datetime.min.time()).replace(hour=18))

    print(f"\n  Report date : {selected_date.strftime('%d-%m-%Y')}")
    print(f"  Window      : {_hr_format_dt_12h(report_start)} → {_hr_format_dt_12h(report_end)}")

    # ── Sleep ────────────────────────────────────────────────────────────────
    sleep_raw_hr, grouped_sleep, sleep_gaps = _hr_build_sleep_block(
        raw_data, sleep_source, tz_obj, report_start, report_end)

    grouped_sleep_processed = None
    if include_sleep_processed_block:
        _, grouped_sleep_processed, _ = _hr_build_sleep_block(
            raw_data, "sleep_processed", tz_obj, report_start, report_end)

    # ── Summary ───────────────────────────────────────────────────────────────
    summary = []
    if sleep_raw_hr:
        stg = {"Deep Sleep": 0, "Light Sleep": 0, "REM": 0, "Awake": 0}
        for r in sleep_raw_hr:
            if r["stage"] in stg:
                stg[r["stage"]] += r["mins"]
        ts = stg["Deep Sleep"] + stg["Light Sleep"] + stg["REM"]
        summary.append({"Metric": "Total Sleep", "Value": f"{ts // 60}h {ts % 60}m"})
        for s, m in stg.items():
            summary.append({"Metric": s, "Value": f"{m // 60}h {m % 60}m"})

    # ── Generic metric helper ─────────────────────────────────────────────────
    def _process_metric(data_keys, val_fields, threshold=15):
        raw = []
        for dk in data_keys:
            for r in raw_data.get(dk, []):
                dt = _hr_parse_to_tz(r.get("date") or r.get("start"), tz_obj)
                if not (dt and report_start <= dt <= report_end):
                    continue
                val = None
                for vf in val_fields:
                    val = r.get(vf)
                    if val is not None:
                        break
                raw.append({"dt": dt, "Date & Time": _hr_format_dt_12h(dt), "Value": val})
        raw.sort(key=lambda x: x["dt"])
        lst  = [{k: v for k, v in r.items() if k != "dt"} for r in raw]
        gps  = _hr_find_gaps(raw, report_start, report_end, threshold)
        return lst, gps

    hr_l,   hr_g   = _process_metric(["heartrate"],
                                      ["bpm", "singleHR", "heartrate", "heartRate"], 10)
    spo2_l, spo2_g = _process_metric(["spo2", "sp02"],
                                      ["automaticSpo2Data", "spo2", "value"], 60)
    temp_l, temp_g = _process_metric(["temperature", "temperature_table"],
                                      ["temperature", "value", "avgTemp", "avg_temp"], 10)

    # ── HRV ──────────────────────────────────────────────────────────────────
    hrv_raw = []
    for dk in ["hrv", "hrv_table"]:
        for r in raw_data.get(dk, []):
            dt = _hr_parse_to_tz(r.get("date"), tz_obj)
            if dt and report_start <= dt <= report_end:
                hrv_raw.append({
                    "dt": dt,
                    "Date & Time": _hr_format_dt_12h(dt),
                    "heartRate":   r.get("heartRate") or r.get("heartrate"),
                    "highBP":      r.get("systolicBP") or r.get("highBP"),
                    "trv":         r.get("hrv") or r.get("trv"),
                    "lowBP":       r.get("diastolicBP") or r.get("lowBP"),
                    "stress":      r.get("stress"),
                })
    hrv_raw.sort(key=lambda x: x["dt"])
    hrv_l = [{k: v for k, v in r.items() if k != "dt"} for r in hrv_raw]
    hrv_g = _hr_find_gaps(hrv_raw, report_start, report_end, 60)

    # ── Export ────────────────────────────────────────────────────────────────
    def _export_styled(filename, data_map, cols_key, style_map=STYLE_MAP):
        print(f"  📝 Exporting: {filename}")
        with pd.ExcelWriter(filename, engine="openpyxl") as writer:
            if filename == report_filename and summary:
                pd.DataFrame(summary).to_excel(writer, sheet_name="Summary", index=False)
            sn = "Report" if filename == report_filename else "Gaps"
            for label_key, df_data in data_map.items():
                start_c = style_map[label_key][cols_key][0]
                pd.DataFrame(df_data).to_excel(writer, sheet_name=sn,
                                               startrow=1, startcol=start_c - 1, index=False)
            wb = writer.book
            sh = wb[sn]
            for label_key, info in style_map.items():
                sc, ec = info[cols_key]
                sh.cell(row=1, column=sc).value = f"--- {label_key} ---"
                fill  = PatternFill(start_color=info["color"], end_color=info["color"], fill_type="solid")
                font  = Font(bold=True, size=12)
                align = Alignment(horizontal="center")
                for c in range(sc, ec + 1):
                    sh.cell(row=1, column=c).fill      = fill
                    sh.cell(row=1, column=c).font      = font
                    sh.cell(row=1, column=c).alignment = align
                    sh.cell(row=2, column=c).font      = Font(bold=True)
                    sh.cell(row=2, column=c).alignment = align
                    sh.column_dimensions[get_column_letter(c)].width = 24

    if include_sleep_processed_block:
        main_data_map = {"SLEEP": grouped_sleep, "SLEEP_PROCESSED": grouped_sleep_processed,
                          "HR": hr_l, "SPO2": spo2_l, "HRV": hrv_l, "TEMP": temp_l}
        _export_styled(report_filename, main_data_map, "main_cols",
                       style_map=STYLE_MAP_WITH_SLEEP_PROCESSED)
    else:
        _export_styled(report_filename,
                       {"SLEEP": grouped_sleep, "HR": hr_l, "SPO2": spo2_l, "HRV": hrv_l, "TEMP": temp_l},
                       "main_cols")
    _export_styled(gaps_filename,
                   {"SLEEP": sleep_gaps, "HR": hr_g, "SPO2": spo2_g, "HRV": hrv_g, "TEMP": temp_g},
                   "gap_cols")

    print(f"  ✅ Main report : {os.path.abspath(report_filename)}")
    print(f"  ✅ Gaps report : {os.path.abspath(gaps_filename)}")

    if return_sleep_summary:
        return _hr_read_report_sleep_times(report_filename)
    return None


# ─────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description="Timon Data Processor")
    parser.add_argument("input", help="Input JSON file (e.g. test_data.json)")
    args = parser.parse_args()

    if not Path(args.input).exists():
        print(f"❌ File not found: {args.input}")
        sys.exit(1)

    # ── Load raw JSON once — shared by all downstream steps ──────────────────
    with open(args.input, "r") as f:
        raw_data = json.load(f)

    has_sleep_processed_in_json = (
        isinstance(raw_data, dict)
        and "sleep_processed" in raw_data
        and isinstance(raw_data["sleep_processed"], list)
        and len(raw_data["sleep_processed"]) > 0
    )

    # ── Steps 1–3 always run ──────────────────────────────────────────────────
    import_json(args.input)
    tz_name, offset_hours, label, tz_obj = ask_timezone()
    convert_dates(tz_name, offset_hours)

    # ── Step 4: sleep processing ──────────────────────────────────────────────
    if has_sleep_processed_in_json:
        print("\n💤 Step 4: sleep_processed data found in source JSON.")
        print("─" * 50)
        sp_path = CORRECTED_DIR / "sleep_processed_timezone_corrected.parquet"
        if sp_path.exists():
            print(f"  ✅ sleep_processed_timezone_corrected files already written by Step 3.")
        else:
            delta = timedelta(hours=offset_hours)
            def _conv(d):
                try:
                    dt = datetime.strptime(str(d), "%Y.%m.%d %H:%M:%S")
                    return (dt + delta).strftime("%Y.%m.%d %H:%M:%S")
                except Exception:
                    return d

            sp_rows = raw_data["sleep_processed"]
            for row in sp_rows:
                if "date" in row:
                    row["date"] = _conv(row["date"])
                if "start" in row:
                    row["start"] = _conv(row["start"])

            df_sp = pd.DataFrame(sp_rows).sort_values("date").reset_index(drop=True)
            df_sp.to_parquet(sp_path, index=False)
            json_sp_path = CORRECTED_DIR / "sleep_processed_timezone_corrected.json"
            with open(json_sp_path, "w") as f:
                json.dump(sp_rows, f, indent=2)
            print(f"  ✅ Written → {sp_path.name}")
            print(f"  ✅ Written → {json_sp_path.name}")

        sleep_processed_available = True

    else:
        print("\n💤 Step 4: No sleep_processed data found in source JSON.")
        print("─" * 50)
        ans = input("  Would you like to generate sleep processed data now? [y/n]: ").strip().lower()
        if ans == "y":
            run_sleep_pipeline(offset_hours)
            sleep_processed_available = True
        else:
            print("  ⏭️  Skipping sleep processing.")
            sleep_processed_available = False

    print("\n" + "═" * 45)
    print("✅ Processing complete!")
    print(f"   Timezone-corrected tables : {CORRECTED_DIR}/")
    if sleep_processed_available:
        print(f"   Sleep processed           : {CORRECTED_DIR}/sleep_processed_timezone_corrected.parquet")
    print("═" * 45)

    # ── Top-level mode menu ───────────────────────────────────────────────────
    while True:
        print("\n" + "═" * 50)
        print("  What would you like to do?")
        print("─" * 50)
        print("  1) Analyse Sleep")
        print("  2) Generate Health Report  (Excel)")
        print("  3) Generate Processed Health Report  (Excel)")
        print("  4) Fetch Sleep Start and End Time")
        print("  0) Exit")
        print()
        mode = input("  Enter your choice [1/2/3/4/0]: ").strip()

        if mode == "1":
            # ── Sleep analysis sub-menu ───────────────────────────────────────
            while True:
                choice = ask_analysis(sleep_processed_available)
                if choice == "1":
                    analyse_sleep_timezone_corrected()
                elif choice == "2" and sleep_processed_available:
                    analyse_sleep_processed_timezone_corrected()
                elif choice == "0":
                    break
                else:
                    print("  ❌ Invalid choice — please enter a valid option.")

        elif mode == "2":
            generate_health_report(raw_data, tz_obj)

        elif mode == "3":
            generate_health_report(raw_data, tz_obj, sleep_source="sleep_processed",
                                    report_filename=REPORT_FILE_PROCESSED,
                                    gaps_filename=GAPS_FILE_PROCESSED)

        elif mode == "4":
            fetch_sleep_start_and_end_time(raw_data, tz_obj, sleep_processed_available)

        elif mode == "0":
            print("\n👋 Goodbye!")
            break
        else:
            print("  ❌ Invalid choice — please enter 1, 2, 3, 4, or 0.")

if __name__ == "__main__":
    main()